from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any
import uuid

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from reports.engineering_report import (
    _build_checks,
    _clean_text,
    _design_basis_rows,
    _element_balance_rows,
    _line_rows,
    _make_profile_chart,
    _result_rows,
    _pump_design_basis_rows,
    _pump_result_rows,
    _pump_npsha_rows,
    _make_system_curve_chart,
)


GENERATED_DIR = Path("generated_reports")
GENERATED_DIR.mkdir(parents=True, exist_ok=True)


def _safe_text(value: Any) -> str:
    return _clean_text(value) if value is not None else ""


def _scenario_payload(item: dict[str, Any], project_title: str) -> dict[str, Any]:
    payload = dict(item)
    payload["project_title"] = f"{project_title} — Scenario: {_safe_text(item.get('name'))}"
    return payload


def _summary_rows(scenarios: list[dict[str, Any]]) -> list[list[str]]:
    rows = [[
        "Scenario",
        "Flow",
        "Total ΔP (bar)",
        "Outlet P (bar a)",
        "Max velocity (m/s)",
        "Warnings",
    ]]

    for scenario in scenarios:
        result = scenario.get("result", {}) or {}
        velocities = [
            float(e["velocity_m_s"])
            for e in (result.get("elements", []) or [])
            if e.get("velocity_m_s") is not None
        ]
        max_velocity = max(velocities) if velocities else None

        def fmt(v: Any, digits: int = 4) -> str:
            if v is None:
                return "—"
            try:
                return f"{float(v):.{digits}f}"
            except Exception:
                return _safe_text(v)

        rows.append([
            _safe_text(scenario.get("name")) or "Unnamed scenario",
            f"{scenario.get('flow_value', '—')} {scenario.get('flow_unit', '')}".strip(),
            fmt(result.get("total_dp_bar"), 4),
            fmt(result.get("outlet_pressure_bar_a"), 3),
            fmt(max_velocity, 3),
            str(len(result.get("warnings", []) or [])),
        ])

    return rows


def _comparison_note(scenarios: list[dict[str, Any]]) -> str:
    successful = [
        s for s in scenarios
        if (s.get("result") or {}).get("total_dp_bar") is not None
    ]
    if not successful:
        return "No calculated scenario results were available for comparison."

    lowest = min(
        successful,
        key=lambda s: float((s.get("result") or {}).get("total_dp_bar"))
    )
    highest = max(
        successful,
        key=lambda s: float((s.get("result") or {}).get("total_dp_bar"))
    )

    return (
        f"Across the saved scenarios, the lowest calculated total pressure drop is "
        f"{float(lowest['result']['total_dp_bar']):.4f} bar for “{_safe_text(lowest.get('name'))}”, "
        f"while the highest is {float(highest['result']['total_dp_bar']):.4f} bar for "
        f"“{_safe_text(highest.get('name'))}”. This comparison is descriptive and does not "
        f"select a preferred design; engineering acceptance criteria remain project-specific."
    )


def _pdf_table(
    story: list[Any],
    rows: list[list[str]],
    widths: list[float],
    cell_style: ParagraphStyle,
    header_style: ParagraphStyle,
    font_size: float = 7.5,
) -> None:
    wrapped = []
    for r_idx, row in enumerate(rows):
        style = header_style if r_idx == 0 else cell_style
        wrapped.append([
            Paragraph(_safe_text(cell).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), style)
            for cell in row
        ])

    table = Table(wrapped, colWidths=widths, repeatRows=1, hAlign="LEFT", splitByRow=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D1D5DB")),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(table)


def create_project_pdf_report(payload: dict[str, Any]) -> Path:
    project_title = _safe_text(payload.get("project_title")) or "Hydraulic Project"
    project_description = _safe_text(payload.get("project_description"))
    scenarios = payload.get("scenarios", []) or []

    if not scenarios:
        raise ValueError("At least one saved scenario is required for a project report.")

    uid = uuid.uuid4().hex[:10]
    base = f"hydraulic_project_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uid}"
    pdf_path = GENERATED_DIR / f"{base}.pdf"

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="ProjectCenterTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=18,
        leading=21,
        spaceAfter=5,
    ))
    cell_style = ParagraphStyle(
        name="ProjectTableCell",
        parent=styles["BodyText"],
        fontSize=7.5,
        leading=9.5,
        wordWrap="CJK",
    )
    header_style = ParagraphStyle(
        name="ProjectTableHeader",
        parent=cell_style,
        fontName="Helvetica-Bold",
    )

    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=A4,
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=f"{project_title} - Hydraulic Scenario Engineering Report",
        author="Process Engineering Workbench",
    )

    story: list[Any] = [
        Paragraph("Process Engineering Workbench", styles["ProjectCenterTitle"]),
        Paragraph("Hydraulic Project & Scenario Engineering Report", styles["Heading2"]),
        Paragraph(f"<b>Project:</b> {_safe_text(project_title)}", styles["BodyText"]),
    ]
    if project_description:
        story.append(Paragraph(f"<b>Description:</b> {_safe_text(project_description)}", styles["BodyText"]))
    story.extend([
        Paragraph(f"<b>Generated:</b> {datetime.now().strftime('%d %B %Y %H:%M')}", styles["BodyText"]),
        Paragraph(f"<b>Saved scenarios included:</b> {len(scenarios)}", styles["BodyText"]),
        Spacer(1, 8),
        Paragraph("1. Scenario Comparison Summary", styles["Heading2"]),
    ])

    _pdf_table(
        story,
        _summary_rows(scenarios),
        [38*mm, 30*mm, 26*mm, 28*mm, 27*mm, 18*mm],
        cell_style,
        header_style,
        7.0,
    )
    story.append(Spacer(1, 7))
    story.append(Paragraph(_safe_text(_comparison_note(scenarios)), styles["BodyText"]))

    for idx, raw_scenario in enumerate(scenarios, start=1):
        scenario = _scenario_payload(raw_scenario, project_title)
        scenario_name = _safe_text(raw_scenario.get("name")) or f"Scenario {idx}"
        description = _safe_text(raw_scenario.get("description"))

        story.append(PageBreak())
        story.append(Paragraph(f"{idx + 1}. Scenario: {scenario_name}", styles["Heading1"]))
        if description:
            story.append(Paragraph(f"<b>Description:</b> {description}", styles["BodyText"]))
        story.append(Spacer(1, 5))

        story.append(Paragraph("Design Basis", styles["Heading2"]))
        _pdf_table(story, _design_basis_rows(scenario), [55*mm, 112*mm], cell_style, header_style)

        story.append(Spacer(1, 7))
        story.append(Paragraph("Process Line Definition", styles["Heading2"]))
        _pdf_table(
            story,
            _line_rows(scenario),
            [8*mm, 31*mm, 37*mm, 63*mm, 28*mm],
            cell_style,
            header_style,
            7.0,
        )

        story.append(Spacer(1, 7))
        story.append(Paragraph("Hydraulic Results", styles["Heading2"]))
        _pdf_table(story, _result_rows(scenario), [65*mm, 102*mm], cell_style, header_style)

        story.append(Spacer(1, 7))
        story.append(Paragraph("Engineering Design Checks", styles["Heading2"]))
        checks = [["Status", "Check", "Value", "Engineering note"]]
        for check in _build_checks(scenario):
            checks.append([check["status"], check["title"], check["value"], check["message"]])
        _pdf_table(
            story,
            checks,
            [18*mm, 36*mm, 31*mm, 82*mm],
            cell_style,
            header_style,
            6.9,
        )

        warnings = (scenario.get("result", {}) or {}).get("warnings", []) or []
        if warnings:
            story.append(Spacer(1, 5))
            story.append(Paragraph("Solver warnings", styles["Heading3"]))
            for warning in warnings:
                story.append(
                    Paragraph(
                        f"<b>{_safe_text(warning.get('code'))}</b>: {_safe_text(warning.get('message'))}",
                        styles["Small"],
                    )
                )

        chart = _make_profile_chart(scenario, f"{base}_scenario_{idx}")
        if chart and chart.exists():
            story.append(Spacer(1, 8))
            story.append(Paragraph("Pressure and Elevation Profile", styles["Heading2"]))
            story.append(Image(str(chart), width=172*mm, height=91*mm))

        story.append(PageBreak())
        story.append(Paragraph(f"{scenario_name} — Element-by-Element Pressure Balance", styles["Heading2"]))
        balance = _element_balance_rows(scenario)
        _pdf_table(
            story,
            balance,
            [7*mm, 25*mm, 39*mm, 16*mm, 16*mm, 14*mm, 17*mm, 20*mm, 13*mm],
            cell_style,
            header_style,
            6.3,
        )

    story.extend([
        PageBreak(),
        Paragraph("Methodology and Limitations", styles["Heading2"]),
        Paragraph(
            "This project report consolidates every saved scenario supplied by the Process Engineering Workbench. "
            "Each scenario is recalculated using the deterministic hydraulic solver at report-generation time so that "
            "the reported results correspond to the saved scenario inputs. Pipe friction uses the engineering engine's "
            "Reynolds-number-dependent friction-factor methods; local losses use K values, including Crane-style "
            "database coefficients where configured. Engineering checks are screening criteria and do not replace "
            "project-, company-, code-, equipment-vendor-, or service-specific design requirements.",
            styles["BodyText"],
        ),
    ])

    doc.build(story)
    return pdf_path


def _docx_table(document: Document, rows: list[list[str]], header: bool = True) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = _safe_text(value)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)
                    if header and r_idx == 0:
                        run.bold = True
    document.add_paragraph()


def create_project_docx_report(payload: dict[str, Any]) -> Path:
    project_title = _safe_text(payload.get("project_title")) or "Hydraulic Project"
    project_description = _safe_text(payload.get("project_description"))
    scenarios = payload.get("scenarios", []) or []

    if not scenarios:
        raise ValueError("At least one saved scenario is required for a project report.")

    uid = uuid.uuid4().hex[:10]
    base = f"hydraulic_project_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uid}"
    docx_path = GENERATED_DIR / f"{base}.docx"

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Process Engineering Workbench")
    run.bold = True
    run.font.size = Pt(18)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("Hydraulic Project & Scenario Engineering Report")
    sr.bold = True
    sr.font.size = Pt(13)

    p = document.add_paragraph()
    p.add_run("Project: ").bold = True
    p.add_run(project_title)
    if project_description:
        p = document.add_paragraph()
        p.add_run("Description: ").bold = True
        p.add_run(project_description)

    p = document.add_paragraph()
    p.add_run("Generated: ").bold = True
    p.add_run(datetime.now().strftime("%d %B %Y %H:%M"))
    p = document.add_paragraph()
    p.add_run("Saved scenarios included: ").bold = True
    p.add_run(str(len(scenarios)))

    document.add_heading("1. Scenario Comparison Summary", level=1)
    _docx_table(document, _summary_rows(scenarios))
    document.add_paragraph(_comparison_note(scenarios))

    for idx, raw_scenario in enumerate(scenarios, start=1):
        scenario = _scenario_payload(raw_scenario, project_title)
        scenario_name = _safe_text(raw_scenario.get("name")) or f"Scenario {idx}"
        description = _safe_text(raw_scenario.get("description"))

        document.add_page_break()
        document.add_heading(f"{idx + 1}. Scenario: {scenario_name}", level=1)
        if description:
            p = document.add_paragraph()
            p.add_run("Description: ").bold = True
            p.add_run(description)

        document.add_heading("Design Basis", level=2)
        _docx_table(document, _design_basis_rows(scenario), header=False)

        document.add_heading("Process Line Definition", level=2)
        _docx_table(document, _line_rows(scenario))

        document.add_heading("Hydraulic Results", level=2)
        _docx_table(document, _result_rows(scenario), header=False)

        document.add_heading("Engineering Design Checks", level=2)
        checks = [["Status", "Check", "Value", "Engineering note"]]
        for check in _build_checks(scenario):
            checks.append([check["status"], check["title"], check["value"], check["message"]])
        _docx_table(document, checks)

        warnings = (scenario.get("result", {}) or {}).get("warnings", []) or []
        if warnings:
            p = document.add_paragraph()
            p.add_run("Solver warnings").bold = True
            for warning in warnings:
                wp = document.add_paragraph(style="List Bullet")
                wp.add_run(f"{_safe_text(warning.get('code'))}: ").bold = True
                wp.add_run(_safe_text(warning.get("message")))

        chart = _make_profile_chart(scenario, f"{base}_scenario_{idx}")
        if chart and chart.exists():
            document.add_heading("Pressure and Elevation Profile", level=2)
            document.add_picture(str(chart), width=Inches(6.7))
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

        document.add_heading("Element-by-Element Pressure Balance", level=2)
        _docx_table(document, _element_balance_rows(scenario))

    document.add_page_break()
    document.add_heading("Methodology and Limitations", level=1)
    document.add_paragraph(
        "This project report consolidates every saved scenario supplied by the Process Engineering Workbench. "
        "Each scenario is recalculated using the deterministic hydraulic solver at report-generation time so that "
        "the reported results correspond to the saved scenario inputs. Pipe friction uses the engineering engine's "
        "Reynolds-number-dependent friction-factor methods; local losses use K values, including Crane-style "
        "database coefficients where configured. Engineering checks are screening criteria and do not replace "
        "project-, company-, code-, equipment-vendor-, or service-specific design requirements."
    )

    document.save(docx_path)
    return docx_path


# ============================================================
# PUMP SIZING PROJECT / MULTI-SCENARIO REPORTS
# ============================================================

def _pump_summary_rows(scenarios: list[dict[str, Any]]) -> list[list[str]]:
    rows = [[
        "Scenario",
        "Flow",
        "Pump head (m)",
        "NPSHa (m)",
        "Pump dP (bar)",
        "Shaft (kW)",
        "Motor min (kW)",
        "Warnings",
    ]]

    for scenario in scenarios:
        result = scenario.get("result", {}) or {}
        duty = result.get("pump_duty", {}) or {}
        warnings = result.get("warnings", []) or []

        def fmt(v: Any, digits: int = 3) -> str:
            if v is None:
                return "-"
            try:
                return f"{float(v):.{digits}f}"
            except Exception:
                return _safe_text(v)

        rows.append([
            _safe_text(scenario.get("name")) or "Unnamed scenario",
            f"{scenario.get('flow_value', '-')} {scenario.get('flow_unit', '')}".strip(),
            fmt(duty.get("required_differential_head_m"), 3),
            fmt((result.get("npsha") or {}).get("available_head_m"), 3),
            fmt(duty.get("required_differential_pressure_bar"), 4),
            fmt(duty.get("shaft_power_kw"), 4),
            fmt(duty.get("minimum_motor_rating_kw"), 4),
            str(len(warnings)),
        ])

    return rows


def _pump_comparison_note(scenarios: list[dict[str, Any]]) -> str:
    successful = [
        s for s in scenarios
        if ((s.get("result") or {}).get("pump_duty") or {}).get("required_differential_head_m") is not None
    ]
    if not successful:
        return "No calculated pump-duty results were available for comparison."

    lowest = min(
        successful,
        key=lambda s: float(((s.get("result") or {}).get("pump_duty") or {}).get("required_differential_head_m")),
    )
    highest = max(
        successful,
        key=lambda s: float(((s.get("result") or {}).get("pump_duty") or {}).get("required_differential_head_m")),
    )

    low_head = float(((lowest.get("result") or {}).get("pump_duty") or {}).get("required_differential_head_m"))
    high_head = float(((highest.get("result") or {}).get("pump_duty") or {}).get("required_differential_head_m"))

    return (
        f"Across the saved pump-sizing scenarios, the lowest calculated required pump head is "
        f"{low_head:.3f} m for \"{_safe_text(lowest.get('name'))}\", while the highest is "
        f"{high_head:.3f} m for \"{_safe_text(highest.get('name'))}\". "
        "This comparison is descriptive and does not constitute pump selection; operating-point, "
        "BEP, NPSH and vendor-curve checks remain separate design stages."
    )


def _pump_warning_text(warning: Any) -> str:
    if isinstance(warning, dict):
        code = _safe_text(warning.get("code"))
        message = _safe_text(warning.get("message"))
        return f"{code}: {message}" if code else message
    return _safe_text(warning)


def create_pump_project_pdf_report(payload: dict[str, Any]) -> Path:
    project_title = _safe_text(payload.get("project_title")) or "Pump Sizing Project"
    project_description = _safe_text(payload.get("project_description"))
    scenarios = payload.get("scenarios", []) or []

    if not scenarios:
        raise ValueError("At least one saved pump-sizing scenario is required for a project report.")

    uid = uuid.uuid4().hex[:10]
    base = f"pump_sizing_project_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uid}"
    pdf_path = GENERATED_DIR / f"{base}.pdf"

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name=f"PumpProjectCenterTitle_{uid}",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=18,
        leading=21,
        spaceAfter=5,
    ))
    cell_style = ParagraphStyle(
        name=f"PumpProjectTableCell_{uid}",
        parent=styles["BodyText"],
        fontSize=7.5,
        leading=9.5,
        wordWrap="CJK",
    )
    header_style = ParagraphStyle(
        name=f"PumpProjectTableHeader_{uid}",
        parent=cell_style,
        fontName="Helvetica-Bold",
    )

    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=A4,
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=f"{project_title} - Pump Sizing Scenario Engineering Report",
        author="Process Engineering Workbench",
    )

    story: list[Any] = [
        Paragraph("Process Engineering Workbench", styles[f"PumpProjectCenterTitle_{uid}"]),
        Paragraph("Pump Sizing Project & Scenario Engineering Report", styles["Heading2"]),
        Paragraph(f"<b>Project:</b> {_safe_text(project_title)}", styles["BodyText"]),
    ]
    if project_description:
        story.append(Paragraph(f"<b>Description:</b> {_safe_text(project_description)}", styles["BodyText"]))

    story.extend([
        Paragraph(f"<b>Generated:</b> {datetime.now().strftime('%d %B %Y %H:%M')}", styles["BodyText"]),
        Paragraph(f"<b>Saved scenarios included:</b> {len(scenarios)}", styles["BodyText"]),
        Spacer(1, 8),
        Paragraph("1. Scenario Comparison Summary", styles["Heading2"]),
    ])

    _pdf_table(
        story,
        _pump_summary_rows(scenarios),
        [30*mm, 22*mm, 22*mm, 20*mm, 22*mm, 20*mm, 22*mm, 14*mm],
        cell_style,
        header_style,
        6.8,
    )
    story.append(Spacer(1, 7))
    story.append(Paragraph(_safe_text(_pump_comparison_note(scenarios)), styles["BodyText"]))

    for idx, raw_scenario in enumerate(scenarios, start=1):
        scenario = _scenario_payload(raw_scenario, project_title)
        scenario_name = _safe_text(raw_scenario.get("name")) or f"Scenario {idx}"
        description = _safe_text(raw_scenario.get("description"))

        story.append(PageBreak())
        story.append(Paragraph(f"{idx + 1}. Scenario: {scenario_name}", styles["Heading1"]))
        if description:
            story.append(Paragraph(f"<b>Description:</b> {description}", styles["BodyText"]))
        story.append(Spacer(1, 5))

        story.append(Paragraph("Pump Design Basis", styles["Heading2"]))
        _pdf_table(story, _pump_design_basis_rows(scenario), [65*mm, 102*mm], cell_style, header_style)

        story.append(Spacer(1, 7))
        story.append(Paragraph("Hydraulic Line Definition", styles["Heading2"]))
        _pdf_table(
            story,
            _line_rows(scenario),
            [8*mm, 31*mm, 37*mm, 63*mm, 28*mm],
            cell_style,
            header_style,
            7.0,
        )

        story.append(Spacer(1, 7))
        story.append(Paragraph("Pump Duty Results", styles["Heading2"]))
        _pdf_table(story, _pump_result_rows(scenario), [70*mm, 97*mm], cell_style, header_style)

        story.append(Spacer(1, 7))
        story.append(Paragraph("NPSHa / Suction Performance", styles["Heading2"]))
        _pdf_table(story, _pump_npsha_rows(scenario), [70*mm, 97*mm], cell_style, header_style)

        chart = _make_system_curve_chart(scenario, f"{base}_scenario_{idx}")
        if chart and chart.exists():
            story.append(Spacer(1, 8))
            story.append(Paragraph("System Curve", styles["Heading2"]))
            story.append(Image(str(chart), width=172*mm, height=91*mm))

        warnings = (scenario.get("result", {}) or {}).get("warnings", []) or []
        if warnings:
            story.append(Spacer(1, 6))
            story.append(Paragraph("Solver warnings", styles["Heading3"]))
            for warning in warnings:
                story.append(Paragraph(_safe_text(_pump_warning_text(warning)), styles["BodyText"]))

        assumptions = (scenario.get("result", {}) or {}).get("assumptions", []) or []
        if assumptions:
            story.append(Spacer(1, 6))
            story.append(Paragraph("Scenario assumptions", styles["Heading3"]))
            for assumption in assumptions:
                story.append(Paragraph(f"- {_safe_text(assumption)}", styles["BodyText"]))

    story.extend([
        PageBreak(),
        Paragraph("Methodology and Limitations", styles["Heading2"]),
        Paragraph(
            "This project report consolidates every saved pump-sizing scenario supplied by the "
            "Process Engineering Workbench. Each scenario is recalculated from its own saved hydraulic "
            "definition and pump boundary conditions before report generation. Required pump head and "
            "power are deterministic hydraulic-duty calculations. Standard motor-size selection, pump "
            "performance curves, operating-point intersection, BEP, NPSHr/NPSHa and vendor selection "
            "checks are outside the current pump-sizing scope unless separately evaluated.",
            styles["BodyText"],
        ),
    ])

    doc.build(story)
    return pdf_path


def create_pump_project_docx_report(payload: dict[str, Any]) -> Path:
    project_title = _safe_text(payload.get("project_title")) or "Pump Sizing Project"
    project_description = _safe_text(payload.get("project_description"))
    scenarios = payload.get("scenarios", []) or []

    if not scenarios:
        raise ValueError("At least one saved pump-sizing scenario is required for a project report.")

    uid = uuid.uuid4().hex[:10]
    base = f"pump_sizing_project_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uid}"
    docx_path = GENERATED_DIR / f"{base}.docx"

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Process Engineering Workbench")
    run.bold = True
    run.font.size = Pt(18)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("Pump Sizing Project & Scenario Engineering Report")
    sr.bold = True
    sr.font.size = Pt(13)

    p = document.add_paragraph()
    p.add_run("Project: ").bold = True
    p.add_run(project_title)
    if project_description:
        p = document.add_paragraph()
        p.add_run("Description: ").bold = True
        p.add_run(project_description)

    p = document.add_paragraph()
    p.add_run("Generated: ").bold = True
    p.add_run(datetime.now().strftime("%d %B %Y %H:%M"))
    p = document.add_paragraph()
    p.add_run("Saved scenarios included: ").bold = True
    p.add_run(str(len(scenarios)))

    document.add_heading("1. Scenario Comparison Summary", level=1)
    _docx_table(document, _pump_summary_rows(scenarios))
    document.add_paragraph(_pump_comparison_note(scenarios))

    for idx, raw_scenario in enumerate(scenarios, start=1):
        scenario = _scenario_payload(raw_scenario, project_title)
        scenario_name = _safe_text(raw_scenario.get("name")) or f"Scenario {idx}"
        description = _safe_text(raw_scenario.get("description"))

        document.add_page_break()
        document.add_heading(f"{idx + 1}. Scenario: {scenario_name}", level=1)
        if description:
            p = document.add_paragraph()
            p.add_run("Description: ").bold = True
            p.add_run(description)

        document.add_heading("Pump Design Basis", level=2)
        _docx_table(document, _pump_design_basis_rows(scenario), header=False)

        document.add_heading("Hydraulic Line Definition", level=2)
        _docx_table(document, _line_rows(scenario))

        document.add_heading("Pump Duty Results", level=2)
        _docx_table(document, _pump_result_rows(scenario), header=False)

        document.add_heading("NPSHa / Suction Performance", level=2)
        _docx_table(document, _pump_npsha_rows(scenario), header=False)

        chart = _make_system_curve_chart(scenario, f"{base}_scenario_{idx}")
        if chart and chart.exists():
            document.add_heading("System Curve", level=2)
            document.add_picture(str(chart), width=Inches(6.7))
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

        warnings = (scenario.get("result", {}) or {}).get("warnings", []) or []
        if warnings:
            document.add_heading("Solver warnings", level=2)
            for warning in warnings:
                document.add_paragraph(_pump_warning_text(warning), style="List Bullet")

        assumptions = (scenario.get("result", {}) or {}).get("assumptions", []) or []
        if assumptions:
            document.add_heading("Scenario assumptions", level=2)
            for assumption in assumptions:
                document.add_paragraph(_safe_text(assumption), style="List Bullet")

    document.add_page_break()
    document.add_heading("Methodology and Limitations", level=1)
    document.add_paragraph(
        "This project report consolidates every saved pump-sizing scenario supplied by the "
        "Process Engineering Workbench. Each scenario is recalculated from its own saved hydraulic "
        "definition and pump boundary conditions before report generation. Required pump head and "
        "power are deterministic hydraulic-duty calculations. Standard motor-size selection, pump "
        "performance curves, operating-point intersection, BEP, NPSHr/NPSHa and vendor selection "
        "checks are outside the current pump-sizing scope unless separately evaluated."
    )

    document.save(docx_path)
    return docx_path

