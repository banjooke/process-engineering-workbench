"""One report-data builder for both formats. Render saved results, never recalculate."""
from datetime import datetime
from io import BytesIO
import re
from xml.sax.saxutils import escape

from reports.report_notice import REPORT_DISCLAIMER

SCREENING_NOTICE = "These checks are engineering screening indicators and do not constitute regulatory approval."
STATUS_LABELS = {"pass": "Within typical range", "warn": "Review", "fail": "Outside selected criterion", "info": "Information only"}
FLOW_METRICS = {"Qavg": "Average flow m³/d", "Qmin": "Minimum flow m³/d", "Qpeak": "Peak flow m³/d", "QmaxDay": "Maximum day m³/d", "QmaxMonth": "Maximum month m³/d", "pf": "Domestic peak factor", "pfOverall": "Overall peak factor", "Qdom": "Domestic contribution m³/d", "Qind": "Industrial contribution m³/d", "Qinf": "Infiltration m³/d", "mmin": "Domestic minimum factor"}
WATER_METRICS = {"domestic": "Domestic demand m³/d", "demand": "Average demand m³/d", "avgProduction": "Average production m³/d", "maxDay": "Maximum day m³/d", "peakHour": "Peak hour m³/d", "plantCapacity": "Plant capacity m³/d"}


def fmt(value):
    if value is None or value == "":
        return "Not available"
    if isinstance(value, (int, float)):
        return f"{value:,.6g}"
    return str(value)


def limit_text(value):
    if isinstance(value, dict):
        return "; ".join(f"{k} {fmt(v)}" for k, v in value.items() if v is not None)
    return f"max {fmt(value)}"


def reference_metadata(draft, standard):
    selected = draft["standardIds"][draft["kind"]]
    jurisdiction = (draft.get("projectCriteria") or {}).get("jurisdiction", "Not specified") if selected == "custom" else "International" if selected.startswith("WHO") else {"ZA": "South Africa", "NG": "Nigeria", "KE": "Kenya", "GH": "Ghana", "EU": "European Union"}.get(selected.split("_")[0], "See reference")
    classification = "User-entered project criteria" if selected == "custom" else "Published guidance or draft reference" if any(v in selected for v in ("WHO", "guideline", "GS175")) else "Regulatory reference - confirm applicability"
    verification = "Provisional — source verification required" if standard["status"] == "to verify" else "User-supplied criteria" if selected == "custom" else "Verified in prototype reference"
    return {"Jurisdiction": jurisdiction, "Classification": classification, "Standard": standard["name"], "Version or edition": standard["version"], "Verification": verification, "Reference checked date": standard["checked"], "Effective date": "Confirm in source; not independently supplied", "Source": standard.get("url", "URL not supplied in reference library"), "Source note": standard.get("note", "")}


def check_source(row):
    source = row["source"]
    if ("Nigeria" in source and "Kenya" in source) or ("Kenya" in source and "Mwache" in source):
        return "Comparative published guidance", "Comparative guidance"
    kind = "Typical engineering range" if any(v in source for v in ("typical", "Typical", "Municipal")) else "Input quality indicator" if any(v in source for v in ("Data consistency", "Input validation", "Measured vs assumed")) else "Published design guidance"
    short = "Metcalf and Eddy" if "M&E" in source else "Qasim and Zhu" if any(v in source for v in ("Q Ex.", "Q Ch.", "Qasim")) else "Ten States guidance" if "Ten States" in source else "Engineering screening"
    return kind, short


def comparison_values(document):
    result, draft = document["result"], document["draft"]
    basis, water = result.get("basis", {}), result.get("water", {})
    flows, loads = basis.get("flows", {}), basis.get("loads", {})
    population = draft["flows" if draft["kind"] == "effluent" else "water"].get("population")
    values = {"Basis type": draft["kind"], "Population people": float(population) if population else None}
    values.update({label: flows.get(key) for key, label in FLOW_METRICS.items()})
    if water:
        values.update({"Average flow m³/d": water["avgProduction"], "Peak flow m³/d": water["peakHour"], "Domestic contribution m³/d": water["domestic"]})
    values.update({f"{key} load kg/d": loads.get(key) for key in ("BOD", "COD", "TSS", "TKN", "TP")})
    status_counts = {status: sum(r["status"] == status for r in result["compliance"]["rows"]) for status in STATUS_LABELS}
    values.update({"Population equivalent": basis.get("PE"), "Plant capacity m³/d": water.get("plantCapacity"), "Plant utilisation": "Not calculated by reference engine", "Selected criterion": result["compliance"]["standard"]["name"], "Reference verification": reference_metadata(draft, result["compliance"]["standard"])["Verification"], "Criterion screening": "; ".join(f"{STATUS_LABELS[status]}: {count}" for status, count in status_counts.items() if count)})
    return values


def build_report_data(project, primary, comparisons, generated_at: datetime):
    scenario, document = primary
    result, draft = document["result"], document["draft"]
    sections = []
    def section(title, headings, rows):
        sections.append({"title": title, "headings": headings, "rows": rows})
    section("Project and scenario", ["Item", "Saved value"], [["Project", project.name], ["Project description", project.description or ""], ["Scenario", scenario.name], ["Scenario description", scenario.description or ""], ["Tool", "Design Basis & Compliance"], ["Report generated UTC", generated_at.isoformat()], ["Calculated UTC", document["calculated_at"]], ["Schema version", document["schema_version"]], ["Engine version", document["engine_version"]], ["Basis", draft["kind"]], ["Mode", draft["mode"]]])
    standard = result["compliance"]["standard"]
    section("Selected reference or project criterion", ["Metadata", "Value"], [[k, v] for k, v in reference_metadata(draft, standard).items()])
    inputs = []
    for group in ("flows", "conc", "loads", "water", "predicted"):
        if group == "water" and draft["kind"] != "drinking" or group in ("flows", "conc", "loads") and draft["kind"] != "effluent":
            continue
        for key, value in draft[group].items():
            unit = "mg/L" if group in ("conc", "predicted") else "factor (peak / average)" if group == "loads" else "people" if key == "population" else "L/c/d" if key == "perCapita" else "h/d" if key == "operatingHours" else "%" if key.endswith("Pct") else "m³/d" if "Flow" in key or key in ("infiltration",) else "factor"
            if group == "predicted":
                unit = "dimensionless" if key == "pH" else "per 100 mL" if key in ("FaecalColi", "Ecoli") else "NTU" if key == "turbidity" else "TCU" if key == "colour" else "mg/L"
            inputs.append([f"{group}: {key}", value or "Not entered", unit, draft["origin"].get(key, "assumed") if group == "conc" else "User entered"])
    inputs.extend([["Domestic peak method", draft["pfMethod"], "method", "User selected"], ["Minimum method", draft["minMethod"], "method", "User selected"], ["Pasted observations", draft["data"] or "None", "one consistent user unit", "User entered"]])
    section("Input data and provenance", ["Parameter", "Value", "Unit", "Origin"], inputs)
    basis = result.get("basis")
    if basis:
        section("Flow results", ["Result", "Value"], [[label, basis["flows"].get(key)] for key, label in FLOW_METRICS.items()])
        section("Pollutant loads", ["Parameter", "Average kg/d", "Peak kg/d", "Origin"], [[key, value, basis["peakLoads"].get(key), draft["origin"].get(key, "assumed")] for key, value in basis["loads"].items()])
        section("Population equivalent", ["Result", "Value", "Basis"], [["Population equivalent", basis["PE"], "PE at 60 g BOD5 per person per day"]])
        section("Domestic peak factor comparison", ["Method", "Factor"], [[r["label"], r["pf"]] for r in basis["pfTable"]])
    if result.get("water"):
        section("Water demand and plant capacity", ["Result", "Value"], [[label, result["water"][key]] for key, label in WATER_METRICS.items()])
    checks = document["checks"]
    section("Design checks", ["Indicator", "Value and unit", "Screening status", "Source"], [[r["item"], f"{fmt(r['value'])} {r['unit']}", STATUS_LABELS[r["status"]], check_source(r)[1]] for r in checks])
    section("Reference screening", ["Parameter", "Selected limit", "Entered treated quality", "Screening status"], [[r["label"], limit_text(r["limit"]), r["value"], STATUS_LABELS[r["status"]] + (" - provisional" if standard["status"] == "to verify" else "")] for r in result["compliance"]["rows"]])
    if result["statistics"]["n"]:
        section("Pasted data statistics", ["Statistic", "Value"], list(result["statistics"].items()))
    comparison = {"scenario_ids": [s.id for s, _ in comparisons], "names": [s.name for s, _ in comparisons], "metrics": []}
    if comparisons:
        values = [comparison_values(d) for _, d in comparisons]
        for key in values[0]:
            nums = [v[key] for v in values if type(v[key]) in (int, float)]
            comparison["metrics"].append({"label": key, "values": [v[key] for v in values], "lowest": min(nums) if nums else None, "highest": max(nums) if nums else None})
        section("Scenario comparison", ["Metric", *comparison["names"]], [[r["label"], *r["values"]] for r in comparison["metrics"]])
        section("Comparison interpretation", ["Note"], [["Read-only comparison of saved compatible scenarios in this project. Numeric minima and maxima are descriptive, not an optimality ranking. Different selected criteria are shown separately."]])
    section("Assumptions and calculation note", ["Method summary"], [[document["calculation_note"]], ["SI units. Prototype example defaults may still be present. Unmeasured influent values remain assumptions. Modes share the same basis equations. Treatment performance and plant utilisation are not calculated in this phase."], [SCREENING_NOTICE], [REPORT_DISCLAIMER]])
    section("Method and references", ["Indicator and classification", "Complete criterion", "Full source and methodology"], [[f"{r['item']} - {check_source(r)[0]}", limit_text({"min": r["range"][0], "max": r["range"][1]}), f"{r['source']} {r['note']}".strip()] for r in checks] + [["Comparative guidance", "Not a governing regulatory requirement", "Mixed-country contextual notes are comparative published guidance. Only the separately selected jurisdiction or user-entered project criteria are compared in reference screening."]])
    return {"title": "Design Basis and Compliance", "subtitle": "Process Engineering Workbench | Water and Wastewater Treatment | BETA prototype", "sections": sections, "comparison": comparison, "screening_notice": SCREENING_NOTICE, "disclaimer": REPORT_DISCLAIMER}


def safe_filename(project, scenario, extension):
    def slug(text):
        return re.sub(r"[^A-Za-z0-9_-]+", "-", str(text)).strip("-_")[:45] or "untitled"
    return f"{slug(project)}_{slug(scenario)}_design-basis.{extension}"


def render_report(data, extension):
    stream = BytesIO()
    if extension == "docx":
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.65)
        section.left_margin = section.right_margin = Inches(0.65)
        document.styles["Normal"].font.name = "Calibri"
        document.styles["Normal"].font.size = Pt(9)
        document.add_paragraph(data["title"], "Title")
        document.add_paragraph(data["subtitle"])
        document.add_paragraph(data["screening_notice"])
        for block in data["sections"]:
            document.add_heading(block["title"], level=1)
            table = document.add_table(rows=1, cols=len(block["headings"]))
            table.style = "Light Shading Accent 1"
            for cell, heading in zip(table.rows[0].cells, block["headings"]):
                cell.text = heading
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in block["rows"]:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = fmt(value)
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(8)
                # Permit natural table page breaks, but keep each data row intact.
                keep = OxmlElement("w:cantSplit")
                table.rows[-1]._tr.get_or_add_trPr().append(keep)
        footer = section.footer.paragraphs[0]
        footer.text = "Process Engineering Workbench | BETA | "
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        document.save(stream)
    else:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
        from matplotlib import get_data_path
        from pathlib import Path
        if "WaterReport" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("WaterReport", str(Path(get_data_path()) / "fonts/ttf/DejaVuSans.ttf")))
        styles = getSampleStyleSheet()
        styles["Heading2"].keepWithNext = True
        styles.add(ParagraphStyle(name="WaterCell", fontName="WaterReport", fontSize=7.5, leading=10, splitLongWords=True))
        styles.add(ParagraphStyle(name="WaterTitle", fontName="WaterReport", fontSize=21, leading=25, spaceAfter=12))
        cell = lambda text: Paragraph(escape(fmt(text)).replace("\n", "<br/>"), styles["WaterCell"])
        story = [Paragraph(data["title"], styles["WaterTitle"]), cell(data["subtitle"]), Spacer(1, 5 * mm), cell(data["screening_notice"])]
        for block in data["sections"]:
            story.extend([Spacer(1, 4 * mm), Paragraph(block["title"], styles["Heading2"])])
            n = len(block["headings"])
            width = 178 * mm
            widths = [width] if n == 1 else [width * .29] + [width * .71 / (n - 1)] * (n - 1)
            table = LongTable([[cell(h) for h in block["headings"]], *[[cell(v) for v in row] for row in block["rows"]]], colWidths=widths, repeatRows=1, splitInRow=1, hAlign="LEFT")
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dceeea")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#cbd5d1")), ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
            story.append(table)
        def footer(canvas, doc):
            canvas.setFont("WaterReport", 7)
            canvas.drawString(16 * mm, 11 * mm, "Process Engineering Workbench | BETA")
            canvas.drawRightString(194 * mm, 11 * mm, str(doc.page))
        SimpleDocTemplate(stream, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=16 * mm, bottomMargin=19 * mm, title=data["title"]).build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()
