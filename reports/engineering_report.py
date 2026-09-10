from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any
import math
import uuid
import textwrap
from xml.sax.saxutils import escape

import matplotlib
matplotlib.use("Agg")  # Non-interactive backend: safe for FastAPI/server-side report generation.
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

GENERATED_DIR = Path(__file__).resolve().parent / "generated"
GENERATED_DIR.mkdir(parents=True, exist_ok=True)


def _safe_float(value: Any, default: float | None = None) -> float | None:
    try:
        if value is None:
            return default
        number = float(value)
        if math.isfinite(number):
            return number
    except (TypeError, ValueError):
        pass
    return default


def _fmt(value: Any, digits: int = 3, suffix: str = "") -> str:
    number = _safe_float(value)
    if number is None:
        return "-"
    return f"{number:.{digits}f}{suffix}"


def _clean_text(value: Any) -> str:
    if value is None:
        return "-"
    return str(value).replace("Δ", "d").replace("–", "-").replace("—", "-")


def _build_checks(payload: dict[str, Any]) -> list[dict[str, str]]:
    result = payload.get("result", {}) or {}
    fluid = payload.get("fluid_config", {}) or {}
    phase_type = str(fluid.get("phase_type", "Liquid"))
    minimum_pressure = _safe_float(result.get("minimum_pressure_bar_a"), 0.0) or 0.0
    inlet_pressure = _safe_float(payload.get("inlet_pressure_bar_a"), 0.0) or 0.0
    total_dp = _safe_float(result.get("total_dp_bar"), 0.0) or 0.0
    maximum_mach = _safe_float(result.get("maximum_mach"), 0.0) or 0.0
    elements = result.get("elements", []) or []

    checks: list[dict[str, str]] = []

    checks.append({
        "status": "WARNING" if minimum_pressure <= 0 else "PASS",
        "title": "Absolute pressure",
        "value": f"{minimum_pressure:.3f} bar(a) minimum",
        "message": (
            "Calculated pressure reaches zero or below. Review the design basis and line model."
            if minimum_pressure <= 0
            else "Calculated pressure remains positive throughout the modeled line."
        ),
    })

    vapor_pressure = _safe_float(fluid.get("vapor_pressure_bar_a"))
    if phase_type.lower() == "liquid" and vapor_pressure is not None:
        margin = minimum_pressure - vapor_pressure
        status = "WARNING" if margin <= 0 else "CAUTION" if margin < 0.2 else "PASS"
        message = (
            "Minimum line pressure is at or below vapor pressure; flashing/cavitation risk requires review."
            if margin <= 0
            else "Pressure margin above vapor pressure is small; review local losses and temperature."
            if margin < 0.2
            else "Minimum pressure remains above vapor pressure with more than 0.2 bar screening margin."
        )
        checks.append({
            "status": status,
            "title": "Vapor-pressure margin",
            "value": f"{margin:.3f} bar",
            "message": message,
        })

    if inlet_pressure > 0:
        ratio = 100.0 * total_dp / inlet_pressure
        status = "WARNING" if ratio > 30 else "CAUTION" if ratio > 10 else "PASS"
        checks.append({
            "status": status,
            "title": "Line pressure-drop ratio",
            "value": f"{ratio:.1f}% of inlet pressure",
            "message": (
                "Pressure loss is a large fraction of inlet pressure. Confirm pipe sizing and available driving pressure."
                if ratio > 30
                else "Pressure loss is material relative to inlet pressure; review against the process pressure budget."
                if ratio > 10
                else "Pressure loss is below 10% of inlet pressure in this screening check."
            ),
        })

    velocities = [
        _safe_float(row.get("velocity_m_s"))
        for row in elements
        if _safe_float(row.get("velocity_m_s")) is not None
    ]
    velocities = [v for v in velocities if v is not None]
    reynolds = [
        _safe_float(row.get("reynolds_number"))
        for row in elements
        if _safe_float(row.get("reynolds_number")) is not None
    ]
    reynolds = [r for r in reynolds if r is not None]

    if velocities:
        vmin, vmax = min(velocities), max(velocities)
        if phase_type.lower() == "liquid":
            status = "WARNING" if vmax > 5 else "CAUTION" if vmax > 3 or vmin < 0.5 else "PASS"
            checks.append({
                "status": status,
                "title": "Liquid velocity screening",
                "value": f"{vmin:.2f}-{vmax:.2f} m/s",
                "message": (
                    "Velocity exceeds 5 m/s. Check erosion, noise, water hammer and service-specific limits."
                    if vmax > 5
                    else "Velocity falls outside the default 0.5-3 m/s screening band. Confirm service-specific limits."
                    if status == "CAUTION"
                    else "Velocity lies within the default 0.5-3 m/s screening band."
                ),
            })
        else:
            status = "WARNING" if maximum_mach >= 0.7 else "CAUTION" if maximum_mach >= 0.3 else "PASS"
            checks.append({
                "status": status,
                "title": "Gas compressibility / Mach screening",
                "value": f"Mach {maximum_mach:.3f} maximum",
                "message": (
                    "High Mach number: compressibility and choking effects require careful review."
                    if maximum_mach >= 0.7
                    else "Compressibility effects are becoming significant; verify the gas-flow method."
                    if maximum_mach >= 0.3
                    else "Maximum Mach number is below 0.3 in this screening check."
                ),
            })

    if reynolds:
        rmin, rmax = min(reynolds), max(reynolds)
        transition = any(2300 <= r <= 4000 for r in reynolds)
        laminar = any(r < 2300 for r in reynolds)
        checks.append({
            "status": "CAUTION" if transition else "PASS",
            "title": "Flow regime",
            "value": f"Re {round(rmin):,}-{round(rmax):,}",
            "message": (
                "At least one section is transitional; friction-factor uncertainty is higher."
                if transition
                else "Laminar flow is present and is handled by the solver."
                if laminar
                else "All hydraulic sections are outside the transitional Reynolds-number range."
            ),
        })

    return checks


def _make_profile_chart(payload: dict[str, Any], base_name: str) -> Path | None:
    profile = (payload.get("result", {}) or {}).get("profile", []) or []
    if not profile:
        return None

    x = [_safe_float(p.get("distance_m"), 0.0) or 0.0 for p in profile]
    pressure = [_safe_float(p.get("pressure_bar_a"), 0.0) or 0.0 for p in profile]
    elevation = [_safe_float(p.get("elevation_m"), 0.0) or 0.0 for p in profile]

    path = GENERATED_DIR / f"{base_name}_profile.png"
    fig, ax1 = plt.subplots(figsize=(9, 4.8))
    ax1.plot(x, pressure, linewidth=2.2, label="Pressure")
    ax1.set_xlabel("Distance (m)")
    ax1.set_ylabel("Pressure bar(a)")
    ax1.grid(True, alpha=0.3)

    ax2 = ax1.twinx()
    ax2.plot(x, elevation, linestyle="--", linewidth=1.8, label="Elevation")
    ax2.set_ylabel("Elevation (m)")

    lines = ax1.get_lines() + ax2.get_lines()
    labels = [line.get_label() for line in lines]
    ax1.legend(lines, labels, loc="best")
    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return path


def _design_basis_rows(payload: dict[str, Any]) -> list[list[str]]:
    fluid = payload.get("fluid_config", {}) or {}
    return [
        ["Project / Study", _clean_text(payload.get("project_title") or "Hydraulic Line Analysis")],
        ["Fluid", _clean_text(fluid.get("fluid"))],
        ["Phase", _clean_text(fluid.get("phase_type"))],
        ["Temperature", _fmt(fluid.get("temperature_c"), 2, " degC")],
        ["Flow rate", f"{_clean_text(payload.get('flow_value'))} {_clean_text(payload.get('flow_unit'))}"],
        ["Inlet pressure", _fmt(payload.get("inlet_pressure_bar_a"), 3, " bar(a)")],
        ["Property source", "Manual" if fluid.get("use_manual_properties") else "CoolProp"],
        ["Density", _fmt(fluid.get("density_kg_m3"), 3, " kg/m3")],
        ["Dynamic viscosity", _fmt(fluid.get("dynamic_viscosity_pa_s"), 7, " Pa.s")],
        ["Vapor pressure", _fmt(fluid.get("vapor_pressure_bar_a"), 4, " bar(a)")],
        ["Molecular weight", _fmt(fluid.get("molecular_weight_kg_kmol"), 3, " kg/kmol")],
        ["Compressibility factor Z", _fmt(fluid.get("compressibility_factor"), 4)],
        ["Cp/Cv", _fmt(fluid.get("gamma"), 4)],
    ]


def _line_rows(payload: dict[str, Any]) -> list[list[str]]:
    rows = [["#", "Type", "Description", "Geometry / basis", "Loss input"]]
    for i, element in enumerate(payload.get("elements", []) or [], start=1):
        etype = _clean_text(element.get("type"))
        desc = _clean_text(element.get("description"))
        if etype == "Pipe":
            basis = f"ID {_fmt(element.get('id_mm'), 2, ' mm')}; L {_fmt(element.get('length_m'), 2, ' m')}; eps {_fmt(element.get('roughness_mm'), 4, ' mm')}"
            loss = f"dz {_fmt(element.get('dz_m'), 2, ' m')}"
        elif etype == "Resistance / Fitting":
            basis = f"ID {_fmt(element.get('id_mm'), 2, ' mm')}; {_clean_text(element.get('database_selection') or element.get('fitting_method') or 'Manual K')}"
            loss = f"K total {_fmt(element.get('k_total'), 4)}"
        elif etype == "Known Equipment ΔP":
            basis = "Known equipment loss"
            loss = f"dP {_fmt(element.get('known_dp_bar'), 4, ' bar')}"
        else:
            basis = "Elevation change"
            loss = f"dz {_fmt(element.get('dz_m'), 2, ' m')}"
        rows.append([str(i), etype, desc, basis, loss])
    return rows


def _result_rows(payload: dict[str, Any]) -> list[list[str]]:
    r = payload.get("result", {}) or {}
    return [
        ["Inlet pressure", _fmt(r.get("inlet_pressure_bar_a"), 3, " bar(a)")],
        ["Outlet pressure", _fmt(r.get("outlet_pressure_bar_a"), 3, " bar(a)")],
        ["Total pressure drop", _fmt(r.get("total_dp_bar"), 4, " bar")],
        ["Resistance pressure drop", _fmt(r.get("resistance_dp_bar"), 4, " bar")],
        ["Static pressure drop", _fmt(r.get("static_dp_bar"), 4, " bar")],
        ["Minimum pressure", _fmt(r.get("minimum_pressure_bar_a"), 3, " bar(a)")],
        ["Mass flow", _fmt(r.get("mass_flow_kg_s"), 4, " kg/s")],
        ["Total line length", _fmt(r.get("total_line_length_m"), 2, " m")],
        ["Net elevation change", _fmt(r.get("net_elevation_change_m"), 2, " m")],
        ["Maximum Mach number", _fmt(r.get("maximum_mach"), 4)],
    ]


def _element_balance_rows(payload: dict[str, Any]) -> list[list[str]]:
    rows = [["#", "Element", "Description", "P in", "P out", "dP", "Velocity", "Re", "f"]]
    for row in (payload.get("result", {}) or {}).get("elements", []) or []:
        dp = sum((_safe_float(row.get(k), 0.0) or 0.0) for k in [
            "pipe_friction_dp_bar", "local_resistance_dp_bar", "equipment_dp_bar", "elevation_dp_bar"
        ])
        rows.append([
            str(row.get("index", "")),
            _clean_text(row.get("element_type")),
            _clean_text(row.get("description")),
            _fmt(row.get("pressure_in_bar_a"), 3),
            _fmt(row.get("pressure_out_bar_a"), 3),
            _fmt(dp, 4),
            _fmt(row.get("velocity_m_s"), 3),
            _fmt(row.get("reynolds_number"), 0),
            _fmt(row.get("friction_factor"), 5),
        ])
    return rows



def _pdf_wrapped_text(value: Any, max_chars: int | None = None) -> str:
    """Return escaped ReportLab Paragraph markup with optional hard line breaks.

    Paragraph normally wraps on spaces, but the explicit wrapping makes table
    behaviour deterministic across ReportLab versions and PDF viewers.
    """
    text = _clean_text(value)
    if max_chars is None or max_chars <= 0:
        return escape(text)

    lines = textwrap.wrap(
        text,
        width=max_chars,
        break_long_words=True,
        break_on_hyphens=True,
        replace_whitespace=True,
        drop_whitespace=True,
    ) or [""]
    return "<br/>".join(escape(line) for line in lines)


def create_pdf_report(payload: dict[str, Any]) -> Path:
    uid = uuid.uuid4().hex[:10]
    base = f"hydraulic_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uid}"
    pdf_path = GENERATED_DIR / f"{base}.pdf"
    chart_path = _make_profile_chart(payload, base)

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CenterTitle", parent=styles["Title"], alignment=TA_CENTER, fontSize=18, leading=22))
    styles.add(ParagraphStyle(name="Small", parent=styles["BodyText"], fontSize=8.5, leading=11))
    styles.add(ParagraphStyle(
        name="TableCell",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8.0,
        leading=10.0,
        spaceAfter=0,
        spaceBefore=0,
        wordWrap="LTR",
    ))
    styles.add(ParagraphStyle(
        name="TableHeader",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=8.0,
        leading=10.0,
        spaceAfter=0,
        spaceBefore=0,
        wordWrap="LTR",
    ))

    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title="Hydraulic Line Analysis Report",
        author="Process Engineering Workbench",
    )

    story: list[Any] = []
    story.append(Paragraph("Process Engineering Workbench", styles["CenterTitle"]))
    story.append(Paragraph("Hydraulic Line Analysis - Engineering Report", styles["Heading2"]))
    story.append(Paragraph(f"Generated: {datetime.now().strftime('%d %B %Y %H:%M')}", styles["BodyText"]))
    story.append(Paragraph("Report layout version: wrapfix-2", styles["Small"]))
    story.append(Spacer(1, 8))

    def add_table(
        rows: list[list[str]],
        widths: list[float] | None = None,
        font_size: float = 8.5,
        hard_wrap_chars: list[int | None] | None = None,
    ):
        # Every cell is a Paragraph.  For critical tables we also insert
        # explicit <br/> line breaks so wrapping is deterministic across
        # ReportLab versions and PDF viewers.
        cell_style = ParagraphStyle(
            name=f"TableCell_{font_size}_{uuid.uuid4().hex[:6]}",
            parent=styles["TableCell"],
            fontSize=font_size,
            leading=max(font_size + 2.0, 9.0),
            wordWrap="CJK",
            splitLongWords=1,
            allowWidows=0,
            allowOrphans=0,
        )
        header_style = ParagraphStyle(
            name=f"TableHeader_{font_size}_{uuid.uuid4().hex[:6]}",
            parent=styles["TableHeader"],
            fontSize=font_size,
            leading=max(font_size + 2.0, 9.0),
            wordWrap="CJK",
            splitLongWords=1,
        )

        wrapped_rows = []
        for row_index, row in enumerate(rows):
            style = header_style if row_index == 0 else cell_style
            wrapped_row = []
            for col_index, cell in enumerate(row):
                max_chars = None
                if hard_wrap_chars and col_index < len(hard_wrap_chars):
                    max_chars = hard_wrap_chars[col_index]
                wrapped_row.append(
                    Paragraph(_pdf_wrapped_text(cell, max_chars), style)
                )
            wrapped_rows.append(wrapped_row)

        table = Table(
            wrapped_rows,
            colWidths=widths,
            repeatRows=1 if len(rows) > 2 else 0,
            hAlign="LEFT",
            splitByRow=1,
        )
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D1D5DB")),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(table)

    story.append(Paragraph("1. Design Basis", styles["Heading2"]))
    add_table(_design_basis_rows(payload), [55 * mm, 110 * mm])
    story.append(Spacer(1, 10))

    story.append(Paragraph("2. Process Line Definition", styles["Heading2"]))
    add_table(_line_rows(payload), [8 * mm, 32 * mm, 38 * mm, 62 * mm, 28 * mm], 7.5)
    story.append(Spacer(1, 10))

    story.append(Paragraph("3. Hydraulic Results", styles["Heading2"]))
    add_table(_result_rows(payload), [65 * mm, 100 * mm])
    story.append(Spacer(1, 10))

    story.append(Paragraph("4. Engineering Design Checks", styles["Heading2"]))
    check_rows = [["Status", "Check", "Value", "Engineering note"]]
    for c in _build_checks(payload):
        check_rows.append([c["status"], c["title"], c["value"], c["message"]])
    add_table(
        check_rows,
        [17 * mm, 36 * mm, 32 * mm, 80 * mm],
        7.5,
        hard_wrap_chars=[10, 24, 22, 48],
    )

    warnings = (payload.get("result", {}) or {}).get("warnings", []) or []
    if warnings:
        story.append(Spacer(1, 6))
        story.append(Paragraph("Solver warnings", styles["Heading3"]))
        for w in warnings:
            story.append(Paragraph(f"<b>{_clean_text(w.get('code'))}</b>: {_clean_text(w.get('message'))}", styles["Small"]))

    if chart_path and chart_path.exists():
        story.append(PageBreak())
        story.append(Paragraph("5. Pressure and Elevation Profile", styles["Heading2"]))
        story.append(Image(str(chart_path), width=175 * mm, height=93 * mm))

    story.append(PageBreak())
    story.append(Paragraph("6. Element-by-Element Pressure Balance", styles["Heading2"]))
    balance = _element_balance_rows(payload)
    t = Table(balance, colWidths=[7*mm, 26*mm, 42*mm, 16*mm, 16*mm, 14*mm, 17*mm, 20*mm, 12*mm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#E5E7EB")),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTSIZE", (0,0), (-1,-1), 6.8),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("GRID", (0,0), (-1,-1), 0.3, colors.HexColor("#D1D5DB")),
        ("LEFTPADDING", (0,0), (-1,-1), 3),
        ("RIGHTPADDING", (0,0), (-1,-1), 3),
    ]))
    story.append(t)

    story.append(Spacer(1, 12))
    story.append(Paragraph("Methodology and limitations", styles["Heading2"]))
    story.append(Paragraph(
        "The report records the design basis supplied to the Process Engineering Workbench and the deterministic hydraulic solver output. "
        "Pipe friction uses Reynolds-number-dependent friction-factor methods implemented in the engineering calculation engine. Local losses are represented using K values, including database-derived Crane-style coefficients where selected. "
        "Engineering checks are screening criteria and do not replace project-, company-, code-, equipment-vendor-, or service-specific design requirements.",
        styles["BodyText"],
    ))

    doc.build(story)
    return pdf_path


def _docx_add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def _docx_add_table(document: Document, rows: list[list[str]], header: bool = False) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = _clean_text(value)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8.5)
                    if header and r_idx == 0:
                        run.bold = True
    document.add_paragraph()


def create_docx_report(payload: dict[str, Any]) -> Path:
    uid = uuid.uuid4().hex[:10]
    base = f"hydraulic_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uid}"
    docx_path = GENERATED_DIR / f"{base}.docx"
    chart_path = _make_profile_chart(payload, base)

    document = Document()
    sec = document.sections[0]
    sec.top_margin = Inches(0.6)
    sec.bottom_margin = Inches(0.6)
    sec.left_margin = Inches(0.65)
    sec.right_margin = Inches(0.65)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Process Engineering Workbench")
    run.bold = True
    run.font.size = Pt(18)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("Hydraulic Line Analysis - Engineering Report")
    sr.bold = True
    sr.font.size = Pt(13)
    date_p = document.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_p.add_run(f"Generated: {datetime.now().strftime('%d %B %Y %H:%M')}")

    _docx_add_heading(document, "1. Design Basis")
    _docx_add_table(document, _design_basis_rows(payload))

    _docx_add_heading(document, "2. Process Line Definition")
    _docx_add_table(document, _line_rows(payload), header=True)

    _docx_add_heading(document, "3. Hydraulic Results")
    _docx_add_table(document, _result_rows(payload))

    _docx_add_heading(document, "4. Engineering Design Checks")
    checks = [["Status", "Check", "Value", "Engineering note"]]
    for c in _build_checks(payload):
        checks.append([c["status"], c["title"], c["value"], c["message"]])
    _docx_add_table(document, checks, header=True)

    warnings = (payload.get("result", {}) or {}).get("warnings", []) or []
    if warnings:
        p = document.add_paragraph()
        p.add_run("Solver warnings").bold = True
        for w in warnings:
            wp = document.add_paragraph(style="List Bullet")
            wp.add_run(f"{_clean_text(w.get('code'))}: ").bold = True
            wp.add_run(_clean_text(w.get("message")))

    if chart_path and chart_path.exists():
        _docx_add_heading(document, "5. Pressure and Elevation Profile")
        document.add_picture(str(chart_path), width=Inches(6.7))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    _docx_add_heading(document, "6. Element-by-Element Pressure Balance")
    _docx_add_table(document, _element_balance_rows(payload), header=True)

    _docx_add_heading(document, "7. Methodology and Limitations")
    document.add_paragraph(
        "The report records the design basis supplied to the Process Engineering Workbench and the deterministic hydraulic solver output. "
        "Pipe friction uses Reynolds-number-dependent friction-factor methods implemented in the engineering calculation engine. "
        "Local losses are represented using K values, including database-derived Crane-style coefficients where selected. "
        "Engineering checks are screening criteria and do not replace project-, company-, code-, equipment-vendor-, or service-specific design requirements."
    )

    document.save(docx_path)
    return docx_path

# ============================================================
# PUMP SIZING REPORTS
# ============================================================

def _pump_design_basis_rows(payload: dict[str, Any]) -> list[list[str]]:
    fluid = payload.get("fluid_config", {}) or {}
    result = payload.get("result", {}) or {}
    return [
        ["Project / Study", _clean_text(payload.get("project_title") or "Pump Sizing Analysis")],
        ["Fluid", _clean_text(fluid.get("fluid"))],
        ["Phase", _clean_text(fluid.get("phase_type"))],
        ["Temperature", _fmt(fluid.get("temperature_c"), 2, " degC")],
        ["Design flow", f"{_clean_text(payload.get('flow_value'))} {_clean_text(payload.get('flow_unit'))}"],
        ["Source pressure", _fmt(payload.get("source_pressure_bar_a"), 3, " bar(a)")],
        ["Destination pressure", _fmt(payload.get("destination_pressure_bar_a"), 3, " bar(a)")],
        ["Reference density", _fmt(result.get("reference_density_kg_m3"), 2, " kg/m3")],
        ["Reference viscosity", _fmt(result.get("reference_dynamic_viscosity_pa_s"), 7, " Pa.s")],
        ["Pump efficiency", _fmt(100 * (_safe_float(payload.get("pump_efficiency"), 0.0) or 0.0), 1, " %")],
        ["Motor margin", _fmt(100 * ((_safe_float(payload.get("motor_margin"), 1.0) or 1.0) - 1.0), 1, " %")],
    ]


def _pump_result_rows(payload: dict[str, Any]) -> list[list[str]]:
    r = payload.get("result", {}) or {}
    system = r.get("system", {}) or {}
    boundary = r.get("boundary", {}) or {}
    duty = r.get("pump_duty", {}) or {}
    return [
        ["Required pump head", _fmt(duty.get("required_differential_head_m"), 3, " m")],
        ["Required pump differential pressure", _fmt(duty.get("required_differential_pressure_bar"), 4, " bar")],
        ["Static head", _fmt(system.get("static_head_m"), 3, " m")],
        ["Resistance / equipment head", _fmt(system.get("resistance_head_m"), 3, " m")],
        ["Boundary pressure head", _fmt(boundary.get("pressure_head_m"), 3, " m")],
        ["Total hydraulic-system head", _fmt(system.get("total_head_m"), 3, " m")],
        ["Hydraulic power", _fmt(duty.get("hydraulic_power_kw"), 4, " kW")],
        ["Shaft power", _fmt(duty.get("shaft_power_kw"), 4, " kW")],
        ["Minimum calculated motor power", _fmt(duty.get("minimum_motor_rating_kw"), 4, " kW")],
        ["Mass flow", _fmt(r.get("mass_flow_kg_s"), 4, " kg/s")],
    ]


def _pump_npsha_rows(payload: dict[str, Any]) -> list[list[str]]:
    result = payload.get("result", {}) or {}
    npsha = result.get("npsha", {}) or {}
    if not npsha:
        return [["NPSHa", "Not available in this calculation"]]

    return [
        ["NPSHa", _fmt(npsha.get("available_head_m"), 3, " m")],
        ["Pump suction pressure", _fmt(npsha.get("suction_pressure_bar_a"), 4, " bar(a)")],
        ["Vapour pressure", _fmt(npsha.get("vapor_pressure_bar_a"), 5, " bar(a)")],
        ["Pressure head above vapour pressure", _fmt(npsha.get("pressure_npsh_m"), 3, " m")],
        ["Suction velocity", _fmt(npsha.get("suction_velocity_m_s"), 3, " m/s")],
        ["Suction velocity head", _fmt(npsha.get("suction_velocity_head_m"), 3, " m")],
        ["Suction resistance head", _fmt(npsha.get("suction_resistance_head_m"), 3, " m")],
        ["Suction static contribution", _fmt(npsha.get("suction_static_head_m"), 3, " m")],
        ["Suction-side elements", _clean_text(npsha.get("suction_element_count"))],
    ]


def _make_system_curve_chart(payload: dict[str, Any], base_name: str) -> Path | None:
    curve = payload.get("system_curve") or {}
    points = curve.get("points", []) or []
    if not points:
        return None
    x = [_safe_float(p.get("flow_value"), 0.0) or 0.0 for p in points]
    total = [_safe_float(p.get("total_head_m"), 0.0) or 0.0 for p in points]
    static = [_safe_float(p.get("static_head_m"), 0.0) or 0.0 for p in points]
    resistance = [_safe_float(p.get("resistance_head_m"), 0.0) or 0.0 for p in points]
    path = GENERATED_DIR / f"{base_name}_system_curve.png"
    fig, ax = plt.subplots(figsize=(9, 4.8))
    ax.plot(x, total, linewidth=2.4, label="Total system head")
    ax.plot(x, static, linestyle="--", linewidth=1.9, label="Static head")
    ax.plot(x, resistance, linestyle=":", linewidth=1.9, label="Resistance / equipment head")
    ax.set_xlabel(f"Flow ({curve.get('flow_unit', payload.get('flow_unit', ''))})")
    ax.set_ylabel("Head (m)")
    ax.grid(True, alpha=0.3)
    ax.legend(loc="best")
    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return path


def create_pump_pdf_report(payload: dict[str, Any]) -> Path:
    uid = uuid.uuid4().hex[:10]
    base = f"pump_sizing_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uid}"
    pdf_path = GENERATED_DIR / f"{base}.pdf"
    curve_path = _make_system_curve_chart(payload, base)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name=f"PumpCenter_{uid}", parent=styles["Title"], alignment=TA_CENTER, fontSize=18, leading=22))
    doc = SimpleDocTemplate(str(pdf_path), pagesize=A4, rightMargin=15*mm, leftMargin=15*mm, topMargin=15*mm, bottomMargin=15*mm,
                            title="Pump Sizing Engineering Report", author="Process Engineering Workbench")
    story: list[Any] = [
        Paragraph("Process Engineering Workbench", styles[f"PumpCenter_{uid}"]),
        Paragraph("Pump Sizing - Engineering Report", styles["Heading2"]),
        Paragraph(f"Generated: {datetime.now().strftime('%d %B %Y %H:%M')}", styles["BodyText"]), Spacer(1, 8),
    ]
    def add(rows, widths=None):
        data=[[Paragraph(escape(_clean_text(c)), styles["BodyText"]) for c in row] for row in rows]
        tb=Table(data, colWidths=widths, hAlign="LEFT")
        tb.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#E5E7EB")),("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#D1D5DB")),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4)]))
        story.append(tb)
    story.append(Paragraph("1. Pump Design Basis", styles["Heading2"])); add(_pump_design_basis_rows(payload), [65*mm,100*mm]); story.append(Spacer(1,8))
    story.append(Paragraph("2. Hydraulic Line Definition", styles["Heading2"])); add(_line_rows(payload), [8*mm,32*mm,38*mm,62*mm,28*mm]); story.append(Spacer(1,8))
    story.append(Paragraph("3. Pump Duty Results", styles["Heading2"])); add(_pump_result_rows(payload), [70*mm,95*mm]); story.append(Spacer(1,8))
    story.append(Paragraph("4. NPSHa / Suction Performance", styles["Heading2"])); add(_pump_npsha_rows(payload), [70*mm,95*mm]); story.append(Spacer(1,8))
    if curve_path and curve_path.exists():
        story.append(Paragraph("5. System Curve", styles["Heading2"])); story.append(Image(str(curve_path), width=170*mm, height=90*mm)); story.append(Spacer(1,8))
    assumptions=(payload.get("result",{}) or {}).get("assumptions",[]) or []
    story.append(Paragraph("6. Assumptions and Limitations", styles["Heading2"]))
    for a in assumptions: story.append(Paragraph(f"• {escape(_clean_text(a))}", styles["BodyText"]))
    story.append(Paragraph("Pump duty is calculated from the deterministic hydraulic model and specified source/destination boundary pressures. The minimum motor power is a calculated requirement before standard motor-size selection. NPSHa is calculated from the defined suction-side hydraulic model. Pump-curve operating point, BEP, NPSHr comparison, required NPSH margin and vendor selection remain outside the current scope until pump performance data are supplied.", styles["BodyText"]))
    doc.build(story); return pdf_path


def create_pump_docx_report(payload: dict[str, Any]) -> Path:
    uid=uuid.uuid4().hex[:10]; base=f"pump_sizing_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uid}"; path=GENERATED_DIR/f"{base}.docx"
    curve_path=_make_system_curve_chart(payload, base)
    document=Document(); sec=document.sections[0]; sec.top_margin=Inches(.6); sec.bottom_margin=Inches(.6); sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
    title=document.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=title.add_run("Process Engineering Workbench"); run.bold=True; run.font.size=Pt(18)
    sub=document.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; sr=sub.add_run("Pump Sizing - Engineering Report"); sr.bold=True; sr.font.size=Pt(13)
    document.add_paragraph(f"Generated: {datetime.now().strftime('%d %B %Y %H:%M')}").alignment=WD_ALIGN_PARAGRAPH.CENTER
    _docx_add_heading(document,"1. Pump Design Basis"); _docx_add_table(document,_pump_design_basis_rows(payload))
    _docx_add_heading(document,"2. Hydraulic Line Definition"); _docx_add_table(document,_line_rows(payload),header=True)
    _docx_add_heading(document,"3. Pump Duty Results"); _docx_add_table(document,_pump_result_rows(payload))
    _docx_add_heading(document,"4. NPSHa / Suction Performance"); _docx_add_table(document,_pump_npsha_rows(payload))
    if curve_path and curve_path.exists():
        _docx_add_heading(document,"5. System Curve"); document.add_picture(str(curve_path),width=Inches(6.7)); document.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    _docx_add_heading(document,"6. Assumptions and Limitations")
    for a in (payload.get("result",{}) or {}).get("assumptions",[]) or []: document.add_paragraph(_clean_text(a),style="List Bullet")
    document.add_paragraph("Pump duty is calculated from the deterministic hydraulic model and specified source/destination boundary pressures. The minimum motor power is a calculated requirement before standard motor-size selection. NPSHa is calculated from the defined suction-side hydraulic model. Pump-curve operating point, BEP, NPSHr comparison, required NPSH margin and vendor selection remain outside the current scope until pump performance data are supplied.")
    document.save(path); return path
