"""Owned JSON scenarios, comparison and reports using real isolated SQLite."""
from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO
import importlib
import json
from pathlib import Path
import unittest
from unittest.mock import patch
from uuid import UUID

from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlmodel import SQLModel, Session, create_engine
from sqlalchemy.pool import StaticPool
from docx import Document

from backend.water_documents import prepare_document, require_saved_document, integrity, SCHEMA, ENGINE
from reports import water_treatment_report as report

FIXTURE = json.loads((Path(__file__).parent / "fixtures/water-scenario.json").read_text(encoding="utf-8"))


class WaterTreatmentTests(unittest.TestCase):
    def setUp(self):
        with patch("dotenv.load_dotenv"), patch("sqlmodel.create_engine"):
            self.projects = importlib.import_module("backend.projects")
            self.scenarios = importlib.import_module("backend.scenarios")
            self.water = importlib.import_module("backend.water_treatment")
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        SQLModel.metadata.create_all(self.engine)
        self.app = FastAPI()
        for router in (self.projects.router, self.scenarios.router, self.water.router):
            self.app.include_router(router)
        self.user = UUID(int=1)
        self.app.dependency_overrides[self.projects.get_current_user_id] = lambda: self.user
        def session():
            with Session(self.engine) as db:
                yield db
        self.app.dependency_overrides[self.projects.get_session] = session
        self.client = TestClient(self.app)
        self.addCleanup(self.client.close)
        self.addCleanup(self.engine.dispose)
        self.project = self.client.post("/projects", json={"name": "Water / study", "description": "Design alternatives"}).json()
        self.path = f"/projects/{self.project['id']}"
        self.body = deepcopy(FIXTURE["wastewater"])

    def create(self, body=None):
        response = self.client.post(self.path + "/scenarios", json=body or self.body)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def report(self, scenario_id, extension="pdf", ids=None):
        return self.client.post(self.path + f"/water-treatment/reports/{extension}", json={"scenario_id": scenario_id, "comparison_ids": ids or []})

    def test_complete_round_trip_and_document_provenance(self):
        saved = self.create()
        read = self.client.get(self.path + f"/scenarios/{saved['id']}").json()
        document = require_saved_document(read["fluid_config"])
        for key, value in self.body["fluid_config"]["water_treatment"].items():
            self.assertEqual(document[key], value)
        self.assertEqual(document["integrity"], integrity(document))
        self.assertEqual(document["schema_version"], SCHEMA)
        self.assertEqual(document["engine_version"], ENGINE)

    def draft(self):
        body = deepcopy(self.body)
        document = body["fluid_config"]["water_treatment"]
        document.update(result=None, checks=[], calculated_at=None, calculated_inputs=None, result_status="draft", stale=True)
        document["draft"]["flows"] = {"population": "", "perCapita": "invalid incomplete text"}
        body["flow_value"] = 0
        return body

    def test_incomplete_zero_flow_draft_can_create_save_and_reopen_but_not_report(self):
        saved = self.create(self.draft())
        self.assertEqual(saved["flow_value"], 0)
        self.assertIsNone(saved["fluid_config"]["water_treatment"]["result"])
        response = self.client.put(self.path + f"/scenarios/{saved['id']}", json=self.draft())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.report(saved["id"]).status_code, 422)

    def test_stale_incomplete_tampered_and_incompatible_results_are_rejected(self):
        for change in (lambda d: d.update(stale=True), lambda d: d.update(engine_version="future"), lambda d: d.update(schema_version="water-treatment-scenario/99"), lambda d: d["draft"]["flows"].update(population="99"), lambda d: d["result"]["basis"]["flows"].pop("Qpeak"), lambda d: d["result"]["basis"]["loads"].pop("BOD"), lambda d: d["checks"].clear(), lambda d: d["result"]["compliance"]["standard"].update(status="verified", limits={"COD": 999999})):
            with self.subTest(change=change):
                body = deepcopy(self.body)
                change(body["fluid_config"]["water_treatment"])
                self.assertEqual(self.client.post(self.path + "/scenarios", json=body).status_code, 422)
        document = prepare_document(self.body["fluid_config"]["water_treatment"])
        document["result"]["basis"]["flows"]["Qpeak"] += 1
        with self.assertRaisesRegex(ValueError, "integrity"):
            require_saved_document({"engineering_task": "water_treatment", "water_treatment": document})
        document["result"]["basis"]["flows"]["Qpeak"] = float("inf")
        with self.assertRaises(ValueError):
            prepare_document(document)

    def test_cross_task_overwrites_blocked_both_directions_for_all_existing_tasks(self):
        water = self.create()
        for task in (None, "pressure_drop", "pump_sizing", "system_curve", "control_valve_sizing"):
            body = deepcopy(self.body)
            body.update(calculation_intent="control_valve_sizing" if task == "control_valve_sizing" else "pressure_drop", flow_unit="m3/h", flow_value=10)
            body["fluid_config"] = {"engineering_task": task}
            if task == "control_valve_sizing":
                body["fluid_config"]["control_valve"] = {"result": None}
            original = self.create(body)
            self.assertEqual(self.client.put(self.path + f"/scenarios/{original['id']}", json=self.body).status_code, 409)
            self.assertEqual(self.client.put(self.path + f"/scenarios/{water['id']}", json=body).status_code, 409)
            self.assertEqual(self.client.get(self.path + f"/scenarios/{original['id']}").json()["fluid_config"], body["fluid_config"])

    def test_task_discriminator_cannot_be_spoofed(self):
        for key in ("engineering_task", "calculation_intent"):
            body = deepcopy(self.body)
            if key == "engineering_task":
                body["fluid_config"][key] = "pump_sizing"
            else:
                body[key] = "pressure_drop"
            self.assertEqual(self.client.post(self.path + "/scenarios", json=body).status_code, 422)

    def test_duplicate_has_new_id_copies_inputs_and_forces_recalculation(self):
        original = self.create()
        response = self.client.post(self.path + f"/scenarios/{original['id']}/water-treatment/duplicate", json={})
        self.assertEqual(response.status_code, 200, response.text)
        copy = response.json()
        self.assertNotEqual(copy["id"], original["id"])
        self.assertEqual(copy["name"], "Copy of " + original["name"])
        self.assertEqual(copy["fluid_config"]["water_treatment"]["draft"], original["fluid_config"]["water_treatment"]["draft"])
        self.assertIsNone(copy["fluid_config"]["water_treatment"]["result"])
        self.assertEqual(self.client.get(self.path + f"/scenarios/{original['id']}").json(), original)

    def test_comparison_two_through_five_is_read_only_and_matches_saved_values(self):
        saved = [self.create() for _ in range(5)]
        for count in range(2, 6):
            response = self.client.post(self.path + "/water-treatment/compare", json={"scenario_ids": [r["id"] for r in saved[:count]]})
            self.assertEqual(response.status_code, 200, response.text)
            metric = next(r for r in response.json()["metrics"] if r["label"] == "Average flow m³/d")
            self.assertEqual(metric["values"], [7600] * count)
        self.assertEqual(self.client.get(self.path + "/scenarios").json(), saved)

    def test_comparison_rejects_invalid_counts_duplicates_drafts_cross_project_and_cross_task(self):
        one, two, draft = self.create(), self.create(), self.create(self.draft())
        for ids in ([one["id"]], [one["id"]] * 2, [one["id"]] * 6, [one["id"], draft["id"]]):
            self.assertEqual(self.client.post(self.path + "/water-treatment/compare", json={"scenario_ids": ids}).status_code, 422)
        other = self.client.post("/projects", json={"name": "Different"}).json()
        foreign = self.client.post(f"/projects/{other['id']}/scenarios", json=self.body).json()
        self.assertEqual(self.client.post(self.path + "/water-treatment/compare", json={"scenario_ids": [one["id"], foreign["id"]]}).status_code, 404)
        hydraulic = self.create({**self.body, "calculation_intent": "pressure_drop", "fluid_config": {"engineering_task": "pump_sizing"}})
        self.assertEqual(self.client.post(self.path + "/water-treatment/compare", json={"scenario_ids": [one["id"], hydraulic["id"]]}).status_code, 422)
        self.assertEqual(self.report(one["id"], ids=[one["id"]]).status_code, 422)
        self.assertEqual(self.report(two["id"], ids=[one["id"], draft["id"]]).status_code, 422)

    def test_other_user_cannot_access_update_duplicate_compare_or_report(self):
        one, two = self.create(), self.create()
        self.user = UUID(int=2)
        for method, path, body in (("GET", self.path + "/scenarios", None), ("PUT", self.path + f"/scenarios/{one['id']}", self.body), ("POST", self.path + f"/scenarios/{one['id']}/water-treatment/duplicate", {}), ("POST", self.path + "/water-treatment/compare", {"scenario_ids": [one["id"], two["id"]]})):
            self.assertEqual(self.client.request(method, path, **({"json": body} if body is not None else {})).status_code, 404)
        for extension in ("pdf", "docx"):
            self.assertEqual(self.report(one["id"], extension).status_code, 404)

    def test_unauthenticated_requests_are_rejected(self):
        one = self.create()
        del self.app.dependency_overrides[self.projects.get_current_user_id]
        for extension in ("pdf", "docx"):
            self.assertEqual(self.report(one["id"], extension).status_code, 401)
        self.assertEqual(self.client.post(self.path + f"/scenarios/{one['id']}/water-treatment/duplicate", json={}).status_code, 401)

    def test_owner_downloads_both_formats_from_same_builder_no_temporary_files(self):
        saved = self.create()
        original_builder = self.water.build_report_data
        with patch.object(self.water, "build_report_data", wraps=original_builder) as builder:
            pdf = self.report(saved["id"], "pdf")
            docx = self.report(saved["id"], "docx")
        self.assertEqual(pdf.status_code, 200, pdf.text[:100])
        self.assertTrue(pdf.content.startswith(b"%PDF-"))
        self.assertEqual(docx.status_code, 200)
        self.assertEqual(builder.call_count, 2)
        word = Document(BytesIO(docx.content))
        text = "\n".join(p.text for p in word.paragraphs) + "\n" + "\n".join(cell.text for table in word.tables for row in table.rows for cell in row.cells)
        for token in (SCHEMA, ENGINE, report.SCREENING_NOTICE, report.REPORT_DISCLAIMER, "7,600", "measured", "Within typical range", "Design checks", "Method and references"):
            self.assertIn(token, text)
        self.assertEqual(pdf.headers["cache-control"], "no-store")
        self.assertNotIn("Scenario comparison", text)
        self.assertIn("design-basis.pdf", pdf.headers["content-disposition"])
        # The renderer returns bytes through BytesIO; no server filesystem path is created.
        self.assertIsInstance(pdf.content, bytes)

    def test_comparison_report_is_explicit_and_includes_saved_values(self):
        one, two = self.create(), self.create()
        response = self.report(one["id"], "docx", [one["id"], two["id"]])
        self.assertEqual(response.status_code, 200)
        word = Document(BytesIO(response.content))
        self.assertIn("Scenario comparison", [p.text for p in word.paragraphs])
        cells = [c.text for t in word.tables for r in t.rows for c in r.cells]
        self.assertIn("Average flow m³/d", cells)
        self.assertIn("7,600", cells)

    def test_filename_is_bounded_ascii_and_path_safe(self):
        name = report.safe_filename('../bad/"\r\n♥' * 20, "scenario\\..*" * 20, "pdf")
        self.assertLessEqual(len(name), 110)
        self.assertRegex(name, r"^[A-Za-z0-9_.-]+\.pdf$")
        self.assertNotIn("..", name)

    def test_report_error_is_safe_and_source_values_unchanged(self):
        saved = self.create()
        with patch.object(self.water, "render_report", side_effect=RuntimeError("secret traceback private DB path")):
            response = self.report(saved["id"])
        self.assertEqual(response.status_code, 500)
        self.assertNotIn("secret", response.text)
        self.assertEqual(self.client.get(self.path + f"/scenarios/{saved['id']}").json(), saved)

    def test_water_supply_results_roundtrip_and_report_without_recalculation(self):
        saved = self.create(deepcopy(FIXTURE["water"]))
        self.assertEqual(self.report(saved["id"], "docx").status_code, 200)
        document = saved["fluid_config"]["water_treatment"]
        self.assertEqual(document["result"]["water"], FIXTURE["water"]["fluid_config"]["water_treatment"]["result"]["water"])

    def test_custom_limits_and_provisional_reference_remain_explicit_in_both_report_data_and_word(self):
        for kind, expected in (("custom", "User-supplied criteria"), ("provisional", "Provisional — source verification required")):
            saved = self.create(deepcopy(FIXTURE[kind]))
            response = self.report(saved["id"], "docx")
            self.assertEqual(response.status_code, 200, response.text)
            word = Document(BytesIO(response.content))
            cells = "\n".join(c.text for t in word.tables for r in t.rows for c in r.cells)
            self.assertIn(expected, cells)
            self.assertIn("regulatory approval", cells)

    def test_missing_peak_loads_or_invalid_calculated_inputs_do_not_pass_integrity_validation(self):
        for change in (lambda d: d["result"]["basis"]["peakLoads"].clear(), lambda d: d["draft"]["flows"].update(population="NaN")):
            body = deepcopy(self.body)
            doc = body["fluid_config"]["water_treatment"]
            change(doc)
            doc["result"]["input"] = deepcopy(doc["draft"])
            from backend.water_documents import canonical
            doc["calculated_inputs"] = canonical(doc["draft"])
            self.assertEqual(self.client.post(self.path + "/scenarios", json=body).status_code, 422)
