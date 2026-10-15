"""Render each supported report format without API, database or network access."""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from reports import engineering_report as current
from reports import project_engineering_report as project
from reports.report_notice import REPORT_DISCLAIMER


class BetaReportTests(unittest.TestCase):
    def test_all_eight_reports_include_notice_and_preserve_payload(self):
        scenario = {
            'name': 'Anonymous evaluation', 'project_title': 'Evaluation',
            'flow_value': 12.345, 'flow_unit': 'm3/h',
            'fluid_config': {'fluid': 'Water', 'density_kg_m3': 997, 'dynamic_viscosity_pa_s': 0.001},
            'elements': [], 'result': {'total_dp_bar': 1.234, 'outlet_pressure_bar_a': 2.345, 'pump_duty': {'required_differential_pressure_bar': 1.234}, 'elements': []},
        }
        for module, prefix, payload in [
            (current, '', scenario), (project, 'project_', {'project_title': 'Evaluation', 'scenarios': [scenario]}),
        ]:
            for pump in ['', 'pump_']:
                for extension in ['pdf', 'docx']:
                    name = f'create_{pump}{prefix}{extension}_report'
                    with self.subTest(builder=name), tempfile.TemporaryDirectory() as directory:
                        inputs = deepcopy(payload)
                        original = deepcopy(inputs)
                        paragraphs = []
                        paragraph_class = module.Paragraph

                        def paragraph(text, *args, **kwargs):
                            paragraphs.append(text)
                            return paragraph_class(text, *args, **kwargs)

                        with patch.object(module, 'GENERATED_DIR', Path(directory)), patch.object(module, 'Paragraph', side_effect=paragraph):
                            path = getattr(module, name)(inputs)
                        self.assertEqual(inputs, original, 'report display changes must not mutate engineering payloads')
                        self.assertTrue(path.exists())
                        self.assertGreater(path.stat().st_size, 1000)
                        if extension == 'pdf':
                            self.assertTrue(path.read_bytes().startswith(b'%PDF-'))
                            self.assertEqual(paragraphs.count(REPORT_DISCLAIMER), 1)
                            self.assertTrue(any('12.345' in str(text) for text in paragraphs), 'engineering flow remains present')
                        else:
                            document = Document(path)
                            self.assertEqual([p.text for p in document.paragraphs].count(REPORT_DISCLAIMER), 1)
                            cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
                            self.assertTrue(any('12.345' in text for text in cells), 'engineering flow remains present')
                            self.assertTrue(any('1.234' in text for text in cells), 'engineering result remains present')


if __name__ == '__main__':
    unittest.main()
